import pandas as pd, re, sys, glob, os, shutil
from collections import Counter
D = r"C:\Users\HAL-USER\Desktop"
GM = os.path.join(D, "GATE MOVE")
# Usage: python yard_gate_summary.py                        -> reports on Desktop (moved into GATE MOVE\<dd.mm.yyyy>)
#        python yard_gate_summary.py 02.10.2026             -> re-run a day (Desktop + GATE MOVE\02.10.2026)
#        python yard_gate_summary.py 01.10.2026 "<folder>"  -> reports in another folder (copied, originals kept)
ARG_DATE = sys.argv[1] if len(sys.argv) > 1 else None
SRC = sys.argv[2] if len(sys.argv) > 2 else None
SEARCH = ([SRC] if SRC else [D]) + ([os.path.join(GM, ARG_DATE)] if ARG_DATE else [])
TPSZ = ['22GP','42GP','45GP','22RE','45RE','22UT','42UT']
MAP = {('20','GP'):'22GP',('40','GP'):'42GP',('40','HC'):'45GP',('20','DC'):'22GP',('40','DC'):'42GP',
       ('40','HQ'):'45GP',('20','RF'):'22RE',('40','RH'):'45RE',('20','RE'):'22RE',
       ('20','OT'):'22UT',('40','OT'):'42UT',('20','UT'):'22UT',('40','UT'):'42UT'}  # open top
CNTR = re.compile(r'^[A-Z]{4}\d{7}$')
# GATE OUT: 16-char booking (HASLS00000000000) = Shipper, anything else (OFFHIRE-TWP-OCT-01, MT-TWP-SEP-28) = Agent
BOOKING = re.compile(r'^[A-Z]{5}\d{11}$')
# GATE IN: haulage / shore job no. (MT-TWP-SEP-28, OFFHIRE-.., I20261725, 26BI01624) = Agent, customer name = Consignee
AGENT = re.compile(r'^(?:[A-Z]+(?: [A-Z]+)?(?:-[A-Z0-9]+){2,}|I\d{8}|\d{2}[A-Z]{2}\d{3,})$')  # also "OFF HIRE-HAST-OCT-01"
TYPES = {'GATE IN': ['Consignee', 'Agent'], 'GATE OUT': ['Shipper', 'Agent']}
def classify(move, cells):
    if move == 'GATE OUT':
        ref = next((x for x in cells if BOOKING.match(x)), None)
        return ('Shipper', ref) if ref else ('Agent', next((x for x in cells if AGENT.match(x)), ''))
    ref = next((x for x in cells if AGENT.match(x)), None)
    return ('Agent', ref) if ref else ('Consignee', '')
def file_date(f):  # dd/mm/yyyy from a report file name (dd.mm.yyyy / dd_mm_yyyy / yyyy-mm-dd)
    b = os.path.basename(f)
    m = re.search(r'(\d{2})[._](\d{2})[._](20\d{2})', b)
    if m: return f'{m[1]}/{m[2]}/{m[3]}'
    m = re.search(r'(20\d{2})-(\d{2})-(\d{2})', b)
    if m: return f'{m[3]}/{m[2]}/{m[1]}'
    m = re.search(r'(?<!\d)(\d{2})[._](\d{2})[._](\d{2})(?!\d)', b)  # dd.mm.yy (e.g. "GATE IN  08.10.26")
    if m: return f'{m[1]}/{m[2]}/20{m[3]}'
def g(p):
    for d in SEARCH:
        m = sorted(glob.glob(os.path.join(d, p)))
        if ARG_DATE: m = [f for f in m if file_date(f) in (None, ARG_DATE.replace('.', '/'))]
        if m: return m[0]
    return None
YARDS = [  # name, loc, file pattern, gate-in sheet, gate-out sheet
 ('SMART BKK','BKK25','Report in-out H-A*.xlsx','GATE IN','GATE OUT'),
 ('BC','BKK27','*BC2_HeungA_Daily_Report*.xls','GATE IN','GATE OUT'),
 ('HAST','LCH27',('GATE IN *.xls','GATE OUT *.xls'),None,None),
 ('CELLO','LCH28','*CELLO*GATE IN_OUT*.xlsx','GATE IN REPORT','GATE OUT REPORT'),
 ('PW','LCHY5','*PW_HeungA_Daily_Report*.xls','GATE IN','GATE OUT'),
]
def sheet(f, name):
    x = pd.ExcelFile(f)
    if name is None: return pd.read_excel(x, x.sheet_names[0], header=None)
    s = next(s for s in x.sheet_names if s.strip().upper() == name)
    return pd.read_excel(x, s, header=None)
def parse(df, move):
    out, bad, cl = Counter(), [], []
    hdr = {}
    for _, r in df.iterrows():
        if not hdr and any(isinstance(x, str) and 'CONTAINER' in x.upper() for x in r)                 and any(isinstance(x, str) and x.strip().upper() == 'SIZE' for x in r):
            hdr = {c: str(x).strip() for c, x in r.items() if pd.notna(x)}; continue
        v = [str(x).strip() for x in r if pd.notna(x) and str(x).strip()]
        idx = next((i for i, x in enumerate(v) if CNTR.match(x)), None)
        if idx is None: continue
        size = typ = None
        for j in range(idx+1, len(v)):
            m = re.match(r"^(20|40|45)\s*'?\s*([A-Z]{2})?$", v[j].replace('.0',''))
            if m:
                size = m.group(1); typ = m.group(2)
                if not typ: typ = next(x for x in v[j+1:] if re.match(r'^[A-Z]{2}$', x))
                break
        k = MAP.get((size, typ))
        cat, ref = classify(move, v)
        if k: out[(cat, k)] += 1; cl.append((v[idx], k, cat, ref, ' | '.join(f'{hdr.get(c, c)}: {x}' for c, x in r.items() if pd.notna(x))))
        elif v[idx] not in {c[0] for c in cl}: bad.append((v[idx], size, typ))  # stray repeat of a listed container: ignore
    return out, bad, cl
rows, issues, yard_cntrs, used_files, skipped = [], [], [], [], []
for name, loc, pat, si, so in YARDS:
    if isinstance(pat, tuple): fi, fo = g(pat[0]), g(pat[1])
    else: fi = fo = g(pat)
    if not fi or not fo:  # this yard's reports have not arrived: summarise the yards that are here
        print(f'SKIP: ไม่พบไฟล์รีพอตของลาน {name} ({pat})'); skipped.append(name); continue
    used_files += [fi, fo]
    for move, f, s in (('GATE IN', fi, si), ('GATE OUT', fo, so)):
        c, bad, cl = parse(sheet(f, s), move)
        yard_cntrs += [(name, loc, move, cn, tp, det, os.path.basename(f), cat, ref) for cn, tp, cat, ref, det in cl]
        for cat in TYPES[move]:
            rows.append({'Yard': name, 'Location': loc, 'Move': move, 'Type': cat, **{t: c.get((cat, t), 0) for t in TPSZ},
                         'Total': sum(n for (ct, _), n in c.items() if ct == cat), 'Source': os.path.basename(f)})
        issues += [(name, move, *b) for b in bad]
if not rows: sys.exit('ไม่พบไฟล์รีพอตของลานใดเลย')
det_all = pd.DataFrame(rows)
print(det_all.drop(columns='Source').to_string(index=False))
# per yard / move totals (used by the system comparison)
res = det_all.groupby(['Yard', 'Location', 'Move'], sort=False).agg({**{t: 'sum' for t in TPSZ}, 'Total': 'sum', 'Source': 'first'}).reset_index()
print('UNMAPPED:', issues)

# report date: from argument, else from the yard file names
dates = Counter(filter(None, map(file_date, dict.fromkeys(used_files))))
if len(dates) > 1: print('WARNING: รีพอตมีหลายวันที่:', dict(dates))
DATE = sys.argv[1].replace('.', '/') if len(sys.argv) > 1 else dates.most_common(1)[0][0]
DAY = DATE.replace('/', '.')

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
wb = Workbook(); ws = wb.active; ws.title = 'Summary'
thin = Side(style='thin', color='999999'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr = ['Yard', 'Location', 'Move', 'Type'] + TPSZ + ['Total']
NC = len(hdr); FIRSTT = 'E'; LASTT = L(4 + len(TPSZ))
ws['A1'] = f'H-A Daily Gate In / Gate Out Summary  —  {DATE}'; ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = ('GATE IN = IEC, IED, IEP | GATE OUT = OEV, OEP (ไม่นับ IER, OER)   ·   GATE IN: Consignee = ชื่อลูกค้า, Agent = เลขงานหัวลาก/ชอร์'
            '   ·   GATE OUT: Shipper = Booking 16 ตัว, Agent = เลขงานหัวลาก/ชอร์'); ws['A2'].font = Font(italic=True, color='666666')
fills = {'GATE IN': 'E2EFDA', 'GATE OUT': 'FCE4D6'}
def style(rng, **kw):
    for c in rng:
        for k, v in kw.items(): setattr(c, k, v)
def block(title, color):
    ws.append([]); ws.append([title]); t = ws.max_row
    ws.merge_cells(start_row=t, start_column=1, end_row=t, end_column=NC)
    style(ws[t], font=Font(bold=True, size=12, color='FFFFFF'), fill=PatternFill('solid', fgColor=color))
    ws.append(hdr); h = ws.max_row
    style(ws[h], font=Font(bold=True), fill=PatternFill('solid', fgColor='D9E1F2'), alignment=Alignment(horizontal='center'))
def body_row(vals, move, bold=False):
    ws.append(vals); n = ws.max_row
    ws.cell(n, NC, f'=SUM({FIRSTT}{n}:{LASTT}{n})')
    style(ws[n], fill=PatternFill('solid', fgColor=fills[move]), font=Font(bold=bold))
    return n
yard_rows = {}   # (move, type) -> sheet rows of every yard
def subtotal_row(label, loc, move, parts):
    n = body_row([label, loc, move, 'TOTAL'], move, bold=True)
    for j in range(5, NC):
        col = L(j); ws.cell(n, j, '=' + '+'.join(f'{col}{k}' for k in parts))
    style(ws[n], fill=PatternFill('solid', fgColor={'GATE IN': 'A9D08E', 'GATE OUT': 'F4B084'}[move]))
# 1) one block per yard: Consignee / Agent / total for GATE IN, Shipper / Agent / total for GATE OUT
for yard, grp in det_all.groupby('Yard', sort=False):
    loc = grp.Location.iloc[0]
    block(f'{yard}  ({loc})', '1F4E78')
    for move, mg in grp.groupby('Move', sort=False):
        parts = []
        for _, r in mg.iterrows():
            n = body_row([r.Yard, r.Location, r.Move, r.Type] + [int(r[t]) for t in TPSZ], r.Move)
            parts.append(n); yard_rows.setdefault((move, r.Type), []).append(n)
        subtotal_row(yard, loc, move, parts)
# 2) total of all yards
block('TOTAL ALL YARDS', 'C00000')
for move, cats in TYPES.items():
    parts = []
    for cat in cats:
        n = body_row(['TOTAL', 'ALL', move, cat], move)
        for j in range(5, NC):
            col = L(j); ws.cell(n, j, '=' + '+'.join(f'{col}{k}' for k in yard_rows[(move, cat)]))
        parts.append(n)
    subtotal_row('TOTAL', 'ALL', move, parts)
for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
    if row[0].value is None: continue
    for c in row:
        c.border = B
        if c.column > 4: c.alignment = Alignment(horizontal='center')
for j in range(1, NC + 1): ws.column_dimensions[L(j)].width = 7
for j, w in ((1, 12), (2, 10), (3, 11), (4, 11), (NC, 8)): ws.column_dimensions[L(j)].width = w
# container list with Consignee / Shipper / Agent split
cs = wb.create_sheet('Container List')
cs.append(['Yard', 'Location', 'Move', 'Type', 'Container No', 'TPSZ', 'Booking / Job ref', 'Yard report detail'])
style(cs[1], font=Font(bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor='1F4E78'))
for name, loc, move, cn, tp, det, f, cat, ref in yard_cntrs:
    cs.append([name, loc, move, cat, cn, tp, ref, det])
for col, w in zip('ABCDEFGH', (11, 9, 10, 11, 14, 7, 20, 120)): cs.column_dimensions[col].width = w
cs.freeze_panes = 'A2'; cs.auto_filter.ref = f'A1:H{cs.max_row}'

# Compare with system export (e.g. 10-2-GATE.xls): GATE IN = IEC/IED/IEP, GATE OUT = OEV/OEP
# system export = any "<m>-<d>-*.xls" (10-2-GATE.xls, 10-2-GATE-UPDATE.xls, 10-1-GAF2.xls ...) whose raw sheet
# has Location / Container No / Move / TPSZ / GateDate for this date; the newest such file wins
SYS_COLS = {'Container No', 'Move', 'Location', 'TPSZ', 'GateDate'}
def raw_sheet(f):  # the sheet holding the raw move list (skips pivot sheets)
    try:
        for sh, d in pd.read_excel(f, sheet_name=None).items():
            if SYS_COLS <= set(map(str, d.columns)): return d
    except Exception:
        return None
_dd, _mm, _yy = DATE.split('/'); GATE_D = f'{_yy}-{_mm}-{_dd}'
on_day = lambda d: pd.to_datetime(d['GateDate'], errors='coerce').dt.strftime('%Y-%m-%d') == GATE_D
_cand = []
for d in SEARCH:
    for f in glob.glob(os.path.join(d, f'{int(_mm)}-{int(_dd)}-*.xls')):
        r = raw_sheet(f)
        if r is not None and on_day(r).any(): _cand.append(f)
SYS = max(_cand, key=os.path.getmtime) if _cand else None
print('SYSTEM FILE:', os.path.basename(SYS) if SYS else 'ไม่พบ (ข้ามการเทียบกับระบบ)')
def read_sys(f):
    d = raw_sheet(f)
    return d[on_day(d)]
ISSUE = {'left_only': 'มีในรีพอตลาน แต่ไม่มีในระบบ', 'right_only': 'มีในระบบ แต่ไม่มีในรีพอตลาน'}
if SYS:
    MOVEG = {'IEC': 'GATE IN', 'IED': 'GATE IN', 'IEP': 'GATE IN', 'OEV': 'GATE OUT', 'OEP': 'GATE OUT'}
    allsys = read_sys(SYS)
    sd = allsys[allsys.Move.isin(MOVEG) & allsys.Location.isin(det_all.Location.unique())]  # only yards reported today
    sysd = pd.DataFrame({'Location': sd.Location, 'Move': sd.Move.map(MOVEG), 'Cntr': sd['Container No'],
                         'TPSZ': sd.TPSZ, 'Code': sd.Move, 'GateTime': sd.GateTime})
    yd = pd.DataFrame(yard_cntrs, columns=['Yard', 'Location', 'Move', 'Cntr', 'TPSZ', 'Detail', 'File', 'Type', 'Ref'])
    # a system export may cover only some locations (e.g. 10-8-LCH27.xls): compare only the locations it contains
    COVERED = set(allsys.Location.dropna())
    not_cov = sorted(set(det_all.Location) - COVERED)
    if not_cov: print('NO SYSTEM DATA (ไม่เทียบ):', not_cov)
    yd = yd[yd.Location.isin(COVERED)]
    res_cmp = res[res.Location.isin(COVERED)]
    m = yd.merge(sysd, on=['Location', 'Move', 'Cntr'], how='outer', suffixes=('_yard', '_sys'), indicator=True)
    c3 = wb.create_sheet('Compare System')
    c3['A1'] = f'Yard report vs System ({os.path.basename(SYS)})  —  {DATE}'; c3['A1'].font = Font(bold=True, size=14)
    c3.append([]); c3.append(['Yard', 'Location', 'Move', 'Source'] + TPSZ + ['Total']); h = c3.max_row
    style(c3[h], font=Font(bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor='1F4E78'), alignment=Alignment(horizontal='center'))
    if not_cov: c3['A2'] = 'ไม่มีข้อมูลระบบ (ไม่ได้เทียบ): ' + ', '.join(not_cov)
    for _, r in res_cmp.iterrows():
        sc = sysd[(sysd.Location == r.Location) & (sysd.Move == r.Move)].TPSZ.value_counts()
        base = c3.max_row + 1
        c3.append([r.Yard, r.Location, r.Move, 'YARD'] + [int(r[t]) for t in TPSZ])
        c3.append([r.Yard, r.Location, r.Move, 'SYSTEM'] + [int(sc.get(t, 0)) for t in TPSZ])
        c3.append([r.Yard, r.Location, r.Move, 'DIFF'])
        for j in range(5, 5 + len(TPSZ)):
            col = L(j); c3.cell(base + 2, j, f'={col}{base}-{col}{base + 1}')
        for n in range(base, base + 3):
            c3.cell(n, 5 + len(TPSZ), f'=SUM(E{n}:{L(4 + len(TPSZ))}{n})')
        ok = int(r.Total) == int(sc.sum())
        style(c3[base + 2], font=Font(bold=True, color='006100' if ok else '9C0006'),
              fill=PatternFill('solid', fgColor='C6EFCE' if ok else 'FFC7CE'))
    for row in c3.iter_rows(min_row=h, max_row=c3.max_row):
        for c in row:
            c.border = B
            if c.column > 4: c.alignment = Alignment(horizontal='center')
    for j in range(1, 6 + len(TPSZ)): c3.column_dimensions[L(j)].width = 7
    for j, w in ((1, 12), (2, 10), (3, 11), (4, 10)): c3.column_dimensions[L(j)].width = w
    c3.freeze_panes = c3.cell(h + 1, 5)
    # container-level differences
    diffs = m[(m._merge != 'both') | (m.TPSZ_yard != m.TPSZ_sys)]
    ynames = {loc: nm for nm, loc, *_ in YARDS}
    DIFF_ROWS = []
    for _, r in diffs.iterrows():
        oth = allsys[allsys['Container No'] == r.Cntr]
        sysmv = ', '.join(f'{a} {b} {int(t):04d}' for a, b, t in zip(oth.Location, oth.Move, oth.GateTime)) or 'ไม่พบในระบบ'
        srow = oth[(oth.Location == r.Location) & (oth.Move == r.Code)] if pd.notna(r.Code) else oth.iloc[0:0]
        sdet = ' | '.join(f'{k}: {v}' for k, v in srow.iloc[0].items() if pd.notna(v)) if len(srow) else ''
        DIFF_ROWS.append({'Yard': ynames.get(r.Location, ''), 'Location': r.Location, 'Move': r.Move, 'Container No': r.Cntr,
                          'TPSZ (Yard)': r.TPSZ_yard, 'TPSZ (System)': r.TPSZ_sys, 'System Code': r.Code,
                          'Issue': ISSUE.get(r._merge, 'TPSZ ไม่ตรงกัน'), 'System moves (this container)': sysmv,
                          'Yard report detail': r.Detail if pd.notna(r.Detail) else '', 'System detail': sdet,
                          'Yard file': r.File if pd.notna(r.File) else ''})
    c4 = wb.create_sheet('Diff Detail')
    c4.append(['Location', 'Move', 'Container No', 'TPSZ (Yard)', 'TPSZ (System)', 'System Code', 'Issue', 'Other system moves of this container'])
    style(c4[1], font=Font(bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor='1F4E78'))
    for d in DIFF_ROWS:
        c4.append([d['Location'], d['Move'], d['Container No'], d['TPSZ (Yard)'], d['TPSZ (System)'], d['System Code'],
                   d['Issue'], d['System moves (this container)']])
    if not DIFF_ROWS: c4.append(['ข้อมูลตรงกันทั้งหมด'])
    for col, w in zip('ABCDEFGH', (10, 11, 14, 11, 13, 12, 32, 50)): c4.column_dimensions[col].width = w
    print('COMPARE: diffs =', len(DIFF_ROWS))
    # ---------- containers missing from the system -> upload rows in the 7-TEMPALTE.xls layout ----------
    # Location | Container No | Move | GateDate (yyyymmdd) | GateTime (HHMM) | BL No (16 chars), all text
    # Move: OUT Shipper / off-hire -> OEV, OUT MT-... job -> OEP, IN with I######## job -> IEP, other IN -> IED
    # BL No: OUT = booking / job no.; IN = import BL, not in the yard reports -> left blank to fill in
    def gate_hhmm(det):
        d = dict(p.split(': ', 1) for p in det.split(' | ') if ': ' in p)
        for k in ('Time', 'TIME'):
            t = re.search(r'(\d{1,2})[:.](\d{2})', str(d.get(k, '')))
            if t: return f'{int(t[1]):02d}{t[2]}'
        t = re.search(r'\b(\d{1,2}):(\d{2})\b', det)
        return f'{int(t[1]):02d}{t[2]}' if t else ''
    def move_code(move, typ, ref):
        if move == 'GATE OUT': return 'OEP' if ref.startswith('MT-') else 'OEV'
        return 'IEP' if re.match(r'^I\d{8}$', ref) else 'IED'
    def bkg_in(r):  # GATE IN BL No = the yard's BKG# (HAST "BKG#", CELLO "BOOKING IN"), else the job no. found in the row
        d = dict(p.split(': ', 1) for p in str(r.Detail).split(' | ') if ': ' in p)
        v = next((d[k].strip() for k in ('BKG#', 'BOOKING IN') if str(d.get(k, '')).strip() not in ('', 'nan')), '')
        return v or str(r.Ref or '')
    miss = m[m._merge == 'left_only']
    MISS_ROWS = [[r.Location, r.Cntr, move_code(r.Move, r.Type, str(r.Ref or '')), GATE_D.replace('-', ''),
                  gate_hhmm(r.Detail), (str(r.Ref or '') if r.Move == 'GATE OUT' else bkg_in(r))[:16]] for _, r in miss.iterrows()]
# Source sheet
s2 = wb.create_sheet('Source Files'); s2.append(['Yard', 'Location', 'Move', 'Source file'])
for _, r in res.iterrows(): s2.append([r.Yard, r.Location, r.Move, r.Source])
if SYS: s2.append(['SYSTEM', '', '', os.path.basename(SYS)])
s2.column_dimensions['D'].width = 60
OUTDIR = os.path.join(GM, DAY); os.makedirs(OUTDIR, exist_ok=True)
LOCKED_OUT = []  # output files open in Excel: skipped with a warning instead of stopping the run
def save_or_skip(path, write):
    try:
        write(); print('saved', path)
    except PermissionError:
        LOCKED_OUT.append(os.path.basename(path)); print('LOCKED (เปิดอยู่ใน Excel ไม่ได้บันทึก):', path)
out = os.path.join(OUTDIR, f'HA Gate In-Out Summary {DAY}.xlsx')
save_or_skip(out, lambda: wb.save(out))
# differences exported to their own workbook
def write_diff(dout, dx):
    with pd.ExcelWriter(dout, engine='openpyxl') as xw:
        dx.to_excel(xw, sheet_name='Diff', index=False, startrow=2)
        w = xw.sheets['Diff']
        w['A1'] = f'Yard report vs System ({os.path.basename(SYS)}) — differences {DATE}  :  {len(dx)} ตู้'
        w['A1'].font = Font(bold=True, size=13)
        for c in w[3]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='C00000'); c.border = B
        for row in w.iter_rows(min_row=4, max_row=w.max_row):
            for c in row: c.border = B; c.alignment = Alignment(vertical='top', wrap_text=c.column >= 9)
        for col, wd in zip('ABCDEFGHIJKL', (10, 9, 10, 14, 10, 12, 11, 28, 26, 60, 60, 40)): w.column_dimensions[col].width = wd
        w.freeze_panes = 'E4'; w.auto_filter.ref = f'A3:L{max(w.max_row, 3)}'
if SYS:
    dx = pd.DataFrame(DIFF_ROWS, columns=['Yard', 'Location', 'Move', 'Container No', 'TPSZ (Yard)', 'TPSZ (System)', 'System Code',
                                          'Issue', 'System moves (this container)', 'Yard report detail', 'System detail', 'Yard file'])
    dout = os.path.join(OUTDIR, f'HA Gate Diff {DAY}.xlsx')
    save_or_skip(dout, lambda: write_diff(dout, dx))
    import xlwt
    MISS = os.path.join(OUTDIR, f'{int(_mm)}-{int(_dd)}-MISSING.xls')
    if MISS_ROWS:
        xb = xlwt.Workbook(); xs = xb.add_sheet('Sheet')
        for j, h in enumerate(['Location', 'Container No', 'Move', 'GateDate', 'GateTime', 'BL No']): xs.write(0, j, h)
        for i, row in enumerate(MISS_ROWS, 1):
            for j, v in enumerate(row): xs.write(i, j, str(v))
        save_or_skip(MISS, lambda: xb.save(MISS))
        print(f'MISSING IN SYSTEM -> {MISS} ({len(MISS_ROWS)} ตู้)')
        for row in MISS_ROWS: print('   ', row)
    elif os.path.exists(MISS):
        os.remove(MISS)  # an earlier run's list; nothing is missing any more
# ---------- running database: GATE MOVE\GATE MOVE DATABASE.xlsx (re-running a day replaces that day) ----------
NAME_KEYS = ['Consignee', 'Customer Name', 'REPO/CONSIGNEE', 'Shippers', 'Shipper Name', 'SHIPPER', 'Shipper']
TIME_KEYS = ['Time', 'TIME']
def fields(det):
    return dict(p.split(': ', 1) for p in det.split(' | ') if ': ' in p)
check = {}
if SYS:
    for _, r in m.iterrows():
        check[(r.Location, r.Move, r.Cntr)] = 'OK' if r._merge == 'both' and r.TPSZ_yard == r.TPSZ_sys \
            else ISSUE.get(r._merge, 'TPSZ ไม่ตรงกัน')
gate_date = pd.to_datetime(DATE, format='%d/%m/%Y')
db_moves = []
for name, loc, move, cn, tp, det, f, cat, ref in yard_cntrs:
    d = fields(det)
    db_moves.append({'Date': gate_date, 'Yard': name, 'Location': loc, 'Move': move, 'Type': cat, 'Container No': cn,
                     'TPSZ': tp, 'Customer': next((d[k].strip() for k in NAME_KEYS if k in d), ''),
                     'Booking / Job ref': ref, 'Gate Time': next((str(d[k]).strip() for k in TIME_KEYS if k in d), ''),
                     'System Check': check.get((loc, move, cn), 'ไม่มีไฟล์ระบบ' if not SYS or loc not in COVERED else ''), 'Source File': f})
db_new = {
    'Container Moves': pd.DataFrame(db_moves),
    'Daily Summary': det_all.drop(columns='Source').assign(Date=gate_date)[['Date', 'Yard', 'Location', 'Move', 'Type'] + TPSZ + ['Total']],
}
if SYS:
    sys_tot = sysd.groupby(['Location', 'Move']).size()
    db_new['System Compare'] = pd.DataFrame([{'Date': gate_date, 'Yard': r.Yard, 'Location': r.Location, 'Move': r.Move,
                                              'Yard Total': int(r.Total), 'System Total': int(sys_tot.get((r.Location, r.Move), 0)),
                                              'Diff': int(r.Total) - int(sys_tot.get((r.Location, r.Move), 0))} for _, r in res_cmp.iterrows()])
    db_new['Diff'] = pd.DataFrame(DIFF_ROWS).drop(columns=['Yard report detail', 'System detail'], errors='ignore').assign(Date=gate_date) \
        if DIFF_ROWS else pd.DataFrame(columns=['Date'])
DB = os.path.join(GM, 'GATE MOVE DATABASE.xlsx')
old = pd.read_excel(DB, sheet_name=None) if os.path.exists(DB) else {}
with pd.ExcelWriter(DB, engine='openpyxl') as xw:
    for sh in dict.fromkeys(list(db_new) + list(old)):
        o = old.get(sh, pd.DataFrame())
        if 'Date' in o.columns: o = o[pd.to_datetime(o['Date']) != gate_date]
        n = db_new.get(sh, pd.DataFrame())
        out_df = pd.concat([x for x in (o, n) if not x.empty], ignore_index=True) if not (o.empty and n.empty) else n
        if 'Date' in out_df.columns:
            out_df['Date'] = pd.to_datetime(out_df['Date'])
            out_df = out_df[['Date'] + [c for c in out_df.columns if c != 'Date']].sort_values('Date', kind='stable')
        out_df.to_excel(xw, sheet_name=sh, index=False)
        w = xw.sheets[sh]
        for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
        for c in w['A'][1:]: c.number_format = 'DD/MM/YYYY'
        for j, col in enumerate(out_df.columns, 1):
            w.column_dimensions[L(j)].width = min(45, max(8, len(str(col)) + 2, *(len(str(v)) + 1 for v in out_df[col].head(300))))
        w.freeze_panes = 'A2'; w.auto_filter.ref = w.dimensions
print('DATABASE updated:', DB, {k: len(v) for k, v in db_new.items()})
# ---------- dashboard: GATE MOVE\HA Gate Dashboard.html built from the whole database ----------
import json, datetime
TPL = os.path.join(GM, 'dashboard_template.html')
if os.path.exists(TPL):
    dbx = pd.read_excel(DB, sheet_name=None)
    cmv, scp = dbx['Container Moves'].fillna(''), dbx.get('System Compare', pd.DataFrame())
    ds = lambda v: pd.to_datetime(v).strftime('%Y-%m-%d')
    payload = {
        'generated': datetime.datetime.now().strftime('%d/%m/%Y %H:%M'),
        'tpsz': TPSZ,
        'yards': [{'name': n, 'loc': l} for n, l, *_ in YARDS],
        'moves': [[ds(r['Date']), r['Yard'], r['Location'], r['Move'], r['Type'], r['Container No'], r['TPSZ'],
                   str(r['Customer']), str(r['Booking / Job ref']), str(r['Gate Time']),
                   '' if r['System Check'] == 'ไม่มีไฟล์ระบบ' else str(r['System Check'])] for _, r in cmv.iterrows()],
        'compare': [[ds(r['Date']), r['Yard'], r['Location'], r['Move'], int(r['Yard Total']), int(r['System Total'])] for _, r in scp.iterrows()],
    }
    html = open(TPL, encoding='utf-8').read()
    data = json.dumps(payload, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    html = re.sub(r'/\*__DATA__\*/\{.*?\};', lambda _: '/*__DATA__*/' + data + ';', html, count=1, flags=re.S)
    DASH = os.path.join(GM, 'HA Gate Dashboard.html')
    open(DASH, 'w', encoding='utf-8').write(html)
    print('DASHBOARD:', DASH)
    # ---------- public build (GitHub Pages, docs\): counts only, no container no. / customer / booking ----------
    DOCS = os.path.join(GM, 'docs'); os.makedirs(DOCS, exist_ok=True)
    pub_check = lambda c: '' if not c else 'OK' if c == 'OK' else 'X'
    pub = dict(payload, public=True,
               moves=[[m[0], m[1], m[2], m[3], m[4], '', m[6], '', '', '', pub_check(m[10])] for m in payload['moves']])
    pdata = json.dumps(pub, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    phtml = re.sub(r'/\*__DATA__\*/\{.*?\};', lambda _: '/*__DATA__*/' + pdata + ';', open(TPL, encoding='utf-8').read(), count=1, flags=re.S)
    open(os.path.join(DOCS, 'dashboard.html'), 'w', encoding='utf-8').write(phtml)
    with pd.ExcelWriter(os.path.join(DOCS, 'gate-move-summary.xlsx'), engine='openpyxl') as xw:
        for sh in ('Daily Summary', 'System Compare'):
            if sh in dbx:
                dbx[sh].to_excel(xw, sheet_name=sh, index=False)
                w = xw.sheets[sh]
                for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
                for c in w['A'][1:]: c.number_format = 'DD/MM/YYYY'
                w.freeze_panes = 'A2'; w.auto_filter.ref = w.dimensions
    print('PUBLIC (docs):', DOCS)
# file all input reports into the day folder
locked = []
for f in dict.fromkeys(used_files + ([SYS] if SYS else [])):
    if os.path.dirname(os.path.abspath(f)) == os.path.abspath(OUTDIR): continue
    dst = os.path.join(OUTDIR, os.path.basename(f))
    if os.path.abspath(os.path.dirname(f)) != os.path.abspath(D):  # from another folder: copy, keep the original
        shutil.copy2(f, dst); print('copied', os.path.basename(f)); continue
    try:
        shutil.move(f, dst); print('moved', os.path.basename(f))
    except PermissionError:  # file is open in Excel: keep a copy in the day folder, original stays on Desktop
        if not os.path.exists(dst): shutil.copy2(f, dst)
        locked.append(os.path.basename(f))
if locked:
    print('COPIED (ไฟล์เปิดอยู่ใน Excel ต้นฉบับยังอยู่บน Desktop ปิดไฟล์แล้วลบทิ้งได้):', locked)
else:
    print('ALL FILES IN', OUTDIR)
if LOCKED_OUT:
    print('NOT SAVED (ปิดไฟล์ใน Excel แล้วรันใหม่):', LOCKED_OUT)
